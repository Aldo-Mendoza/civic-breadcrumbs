from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("Civic_Breadcrumbs_5_Minute_Consulting_Pitch.docx")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = "CIVIC BREADCRUMBS   FIVE MINUTE PITCH"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.runs[0]
    hr.font.name = "Aptos"
    hr.font.size = Pt(8)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor(48, 92, 92)

    footer = section.footer
    footer.is_linked_to_previous = False
    add_page_number(footer.paragraphs[0])


def add_demo_cue(doc, text):
    p = doc.add_paragraph(style="Demo Cue")
    r = p.add_run("Live demo cue  ")
    r.bold = True
    r.font.color.rgb = RGBColor(48, 92, 92)
    r = p.add_run(text)
    r.italic = True
    return p


def add_speech(doc, text):
    p = doc.add_paragraph(style="Speech")
    p.add_run(text)
    return p


def add_speaker_page(doc, part, name, role, timing, paragraphs, cues, new_page=True):
    if new_page:
        doc.add_page_break()

    kicker = doc.add_paragraph(style="Kicker")
    kicker.add_run(f"PART {part}   {timing}")

    title = doc.add_paragraph(style="Speaker Heading")
    title.add_run(f"{name}  {role}")

    cue_map = {index: cue for index, cue in cues}
    for index, text in enumerate(paragraphs):
        if index in cue_map:
            add_demo_cue(doc, cue_map[index])
        add_speech(doc, text)


doc = Document()
configure_section(doc.sections[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11.5)
normal.font.color.rgb = RGBColor(24, 28, 30)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(6)

for style_name, size, bold, before, after in (
    ("Speaker Heading", 19, True, 0, 12),
    ("Kicker", 9, True, 0, 5),
    ("Speech", 11.5, False, 0, 8),
    ("Demo Cue", 9.5, False, 2, 7),
):
    style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Aptos"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_together = True

styles["Kicker"].font.color.rgb = RGBColor(48, 92, 92)
styles["Kicker"].paragraph_format.keep_with_next = True
styles["Speaker Heading"].paragraph_format.keep_with_next = True
styles["Demo Cue"].paragraph_format.left_indent = Inches(0.22)
styles["Demo Cue"].paragraph_format.right_indent = Inches(0.22)

# First page title and purpose
p = doc.add_paragraph(style="Title")
p.add_run("Civic Breadcrumbs Five Minute Consulting Pitch")

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(14)
r = meta.add_run("Four speakers   Five minutes   Live product demonstration")
r.font.name = "Aptos"
r.font.size = Pt(10.5)
r.font.bold = True
r.font.color.rgb = RGBColor(48, 92, 92)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(16)
intro.add_run(
    "This script presents Civic Breadcrumbs as a consulting solution for continuity across fragmented public services. "
    "Each speaker explains the value of the solution, one implementation challenge, how the team addressed it, and one lesson learned."
)

add_speaker_page(
    doc,
    1,
    "Eduardo",
    "Main Idea and Implementation",
    "0:00 to 1:15",
    [
        "Good afternoon. We are presenting Civic Breadcrumbs, a consulting solution built around a common service failure: people must reconstruct their story every time they reach a new public institution.",
        "Consider Amina, an international student in Ottawa. She submits a study permit extension, receives an email, contacts her university and calls IRCC. Each organization remembers only its part. Amina must connect the entire process herself.",
        "Civic Breadcrumbs gives her a citizen-owned Journey. She records each interaction in plain language, and the application builds a timeline showing where she left off, which organization is responsible and what instruction she last received. It recognizes phone calls, emails, letters, websites and in-person visits, and it labels the source of every record.",
        "Our challenge was making the product useful without creating another complicated government form. We solved it by letting people write naturally and requiring confirmation before the interpretation becomes evidence. I learned that inclusive design begins by reducing what people must remember and repeat during a stressful process.",
    ],
    [(2, "Open the seeded Study Permit Extension Journey and point to the timeline and You are here panel")],
    new_page=False,
)

add_speaker_page(
    doc,
    2,
    "Aldo",
    "Technical Implementation",
    "1:15 to 2:30",
    [
        "We designed Civic Breadcrumbs as a modular and dependable web application. The backend uses Python, Django and Django REST Framework. Django manages the data, ownership rules, authentication and API. Separate modules handle Journeys, breadcrumbs, organizations, official sources and AI-assisted interpretation. The application uses SQLite locally and supports PostgreSQL in production.",
        "When a citizen enters a sentence, the system turns it into structured information but does not save it immediately. The citizen reviews what the system understood and confirms or edits it. Only confirmed information becomes evidence.",
        "The main technical challenge was using AI without making the service depend on AI. We created two interchangeable engines. Gemini can interpret natural language, while a deterministic rule-based engine preserves the essential workflow during an outage or quota limit. Each action can make at most one AI request, and ordinary reads make none.",
        "I learned that responsible AI architecture means choosing the few tasks where a model adds value and keeping the rest predictable, affordable and testable.",
    ],
    [(1, "Enter: I called IRCC today. They said it is still processing and told me not to submit another application")],
)

add_speaker_page(
    doc,
    3,
    "Juan Pablo",
    "Impact and Public Value",
    "2:30 to 3:40",
    [
        "The value of Civic Breadcrumbs is continuity. People navigating public services may move between federal, provincial, municipal and institutional organizations. Newcomers, international students and older adults can face additional language, process or digital barriers.",
        "After we confirm this interaction, the application updates Amina's Journey. Notice that it does not claim to know her current immigration status. It says her last recorded interaction reported that the application was processing. That distinction protects trust because the system organizes evidence without pretending to speak for the government.",
        "The I am stuck view explains what happened, what remains unresolved, the latest instruction and the responsible organization. Hand me off creates a concise summary that Amina can take to a university adviser or government employee, so she does not have to reconstruct the case from memory.",
        "Our challenge was helping citizens without giving legal advice or inventing requirements. We solved it with verified official links, visible source labels and clear product boundaries. I learned that bringing people closer to government can begin by helping them navigate existing services with better information and confidence.",
    ],
    [
        (1, "Confirm the interpreted breadcrumb and show the updated You are here panel"),
        (2, "Open I am stuck and then Hand me off"),
    ],
)

add_speaker_page(
    doc,
    4,
    "Cesar",
    "Technical Decisions and Assurance",
    "3:40 to 5:00",
    [
        "Our hardest technical question was trust. How could we use AI without allowing generated text to be mistaken for official evidence?",
        "We separated the citizen's confirmed record, suggested guidance and AI-generated wording. The current state is calculated deterministically from confirmed breadcrumbs only. AI-generated content cannot affect that calculation, and a generated handoff is never stored as evidence.",
        "We also designed for real operating conditions. If Gemini fails, reaches its quota or loses connectivity, the rule-based engine takes over immediately. Users can still record events, view timelines, identify organizations and generate handoffs. The data model minimizes sensitive information and has no dedicated fields for passport numbers, Social Insurance Numbers, addresses or financial details.",
        "The trade-off was choosing reliability over a more impressive but fragile AI demonstration. We addressed it with a modular architecture, explicit confirmation, source labels, access controls, idempotent operations and automated tests for the core workflow and safety rules. I learned that technical quality is visible when dependencies fail and the product still protects the user.",
        "As consultants, we recommend piloting Civic Breadcrumbs with international students and newcomer-support organizations in Ottawa, measuring repeated explanations and missed follow-ups, and then expanding carefully to other public-service journeys. Government should not make people start over.",
    ],
    [],
)

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "Civic Breadcrumbs Five Minute Consulting Pitch"
doc.core_properties.subject = "Four speaker hackathon presentation script"
doc.core_properties.author = "Civic Breadcrumbs Team"
doc.save(OUTPUT)
print(OUTPUT.resolve())
